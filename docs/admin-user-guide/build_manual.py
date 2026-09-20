"""Dựng Word từ Markdown đã duyệt; không gọi ứng dụng hay dịch vụ mạng."""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent


def inline(paragraph, text):
    # Chỉ hỗ trợ tập Markdown nhỏ dùng trong tài liệu này.
    text = text.replace(r"\*", "\uE000")
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(value.replace("\uE000", "*"))
        run.bold = bold


def field(paragraph, instruction, display=None):
    run = paragraph.add_run()
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    start.set(qn("w:dirty"), "true")
    run._r.append(start)
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    paragraph.add_run()._r.append(code)
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separator)
    if display:
        paragraph.add_run(display)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def new_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1.")):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:val"), value)
        level.append(el)
    pp = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "360")
    indent.set(qn("w:hanging"), "360")
    pp.append(indent)
    level.append(pp)
    abstract.append(level)
    numbering.append(abstract)
    instance = numbering.add_num(abstract_id)
    return instance.numId


def table(doc, rows):
    cells = [[v.strip() for v in row.strip().strip("|").split("|")] for row in rows]
    cells = [row for row in cells if not all(re.fullmatch(r":?-+:?", c) for c in row)]
    if not cells:
        return
    t = doc.add_table(rows=0, cols=len(cells[0]))
    t.style = "Table Grid"
    t.autofit = False
    widths = {5: [3.2, 3.0, 2.0, 4.0, 4.8], 4: [4.0, 3.2, 5.8, 4.0], 3: [4.0, 6.0, 7.0]}.get(len(cells[0]))
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
    for index, values in enumerate(cells):
        row = t.add_row()
        for n, value in enumerate(values):
            cell = row.cells[n]
            if widths:
                cell.width = Cm(widths[n])
            p = cell.paragraphs[0]
            inline(p, value)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(4)
            for r in p.runs:
                r.font.size = Pt(10)
                if index == 0:
                    r.bold = True
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7EEF3")
                cell._tc.get_or_add_tcPr().append(shade)
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)
    sec.footer_distance = Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet", "List Number"]:
        s = doc.styles[name]
        s.font.name = "Arial"
        fonts = s.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attr), "Arial")
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "vi-VN")
        s.element.get_or_add_rPr().append(lang)
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("183F55")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Trang ")
    field(footer, " PAGE ")
    for r in footer.runs:
        r.font.size = Pt(9)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    core = doc.core_properties
    core.title = "Hướng dẫn sử dụng trang quản trị website du lịch"
    core.subject = "Travel Admin"
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.version = "0.1"
    lines = (ROOT / "huong-dan-su-dung-admin.md").read_text(encoding="utf-8").splitlines()
    i, in_numbered, num_id = 0, False, None
    image_count = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            table(doc, rows)
            in_numbered = False
            continue
        m = re.match(r"(!\[([^\]]+)\]\(([^)]+)\))$", line)
        if m:
            relative = Path(m.group(3))
            if relative.is_absolute() or ".." in relative.parts:
                raise ValueError("Ảnh phải là đường dẫn tương đối trong thư mục tài liệu")
            path = ROOT / relative
            if not path.exists():
                raise FileNotFoundError(path)
            from PIL import Image
            with Image.open(path) as im:
                width, height = im.size
            display_width = min(17, 19 * width / height)
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            shape = p.add_run().add_picture(str(path), width=Cm(display_width))
            shape._inline.docPr.set("descr", m.group(2))
            doc.add_paragraph(m.group(2), "Caption")
            image_count += 1
        elif line.startswith("# "):
            p = doc.add_paragraph(line[2:], "Title")
            p.paragraph_format.space_before = Cm(2)
            p.paragraph_format.space_after = Pt(20)
        elif line == "## Mục lục":
            doc.add_page_break()
            doc.add_paragraph("Mục lục", "Title")
            field(doc.add_paragraph(), ' TOC \\o "1-2" \\h \\z \\u ', "Cập nhật mục lục trong Word để hiển thị tiêu đề và số trang.")
        elif line.startswith("## "):
            if line.startswith("## 1 "):
                doc.add_page_break()
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        elif re.match(r"^\d+\. ", line):
            if not in_numbered:
                num_id = new_numbering(doc)
            p = doc.add_paragraph()
            np = p._p.get_or_add_pPr().get_or_add_numPr()
            np.get_or_add_ilvl().val = 0
            np.get_or_add_numId().val = num_id
            inline(p, re.sub(r"^\d+\. ", "", line))
            in_numbered = True
            i += 1
            continue
        elif line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            p = doc.add_paragraph()
            inline(p, line)
            if line.startswith("**Hình minh họa"):
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor.from_string("596775")
        in_numbered = False
        i += 1
    output = ROOT / "huong-dan-su-dung-admin.docx"
    doc.save(output)
    print(f"Đã tạo {output.name}; số ảnh nhúng: {image_count}")


if __name__ == "__main__":
    argparse.ArgumentParser(description=__doc__).parse_args()
    build()
